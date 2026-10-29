#!/usr/bin/env python3
"""
Scrape Amazon US product links into Create.xlsx for create_listing.py.

Urls.xlsx columns:
  amazon_url, store_name, sku, upc
  status and free_shipping are optional. Blank means Draft and False.
  quantity, Category, and Subcategory are scraped from the product page.

The scraper sets the Amazon delivery ZIP to 10001, keeps up to 6
photos on any background, skips shots with a person, crops logos on the
edge, and stops when every Amazon link is in Create.xlsx.
Put AliExpress links in scrape_aliexpress/Urls.xlsx and run that bot separately.
Check Create.xlsx, then run:

  python create_listing/create_listing.py create_listing/Create.xlsx
"""
from __future__ import annotations

import io
import json
import re
import time
from pathlib import Path
from typing import Any, Optional
import sys
from urllib.parse import quote

import cv2
import numpy as np

import requests
from openpyxl import Workbook, load_workbook
from PIL import Image
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

BOT_DIR = Path(__file__).resolve().parent
ROOT = BOT_DIR.parent
sys.path.insert(0, str(ROOT / "create_listing"))

from create_listing import local_chromedriver, log

URLS_PATH = BOT_DIR / "Urls.xlsx"
CREATE_PATH = ROOT / "create_listing" / "Create.xlsx"
PHOTO_DIR = ROOT / "output" / "amazon_photos"
ZIP_CODE = "10001"
MAX_IMAGES = 6
IMAGE_SIZE = 1600

URL_COLUMNS = ("amazon_url", "store_name", "quantity", "category", "subcategory", "status", "free_shipping")
CREATE_COLUMNS = (
    "store_name",
    "title",
    "description",
    "remove_branding",
    "price",
    "quantity",
    "Category",
    "Subcategory",
    "sku",
    "photos",
    "status",
    "upc",
    "free_shipping",
    "amazon_url",
)
ASIN_RE = re.compile(r"/(?:dp|gp/product)/([A-Z0-9]{10})", re.I)
HIRES_RE = re.compile(r'"hiRes"\s*:\s*"(https://[^"]+)"')
LARGE_RE = re.compile(r'"large"\s*:\s*"(https://[^"]+)"')


def header_key(value: Any) -> str:
    return re.sub(r"[\s_]+", "", str(value or "").strip().lower())


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none"}:
        return ""
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def ensure_urls_file(path: Path = URLS_PATH) -> None:
    if path.is_file():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "urls"
    ws.append(["amazon_url", "store_name", "sku", "upc", "status", "free_shipping"])
    wb.save(path)
    raise SystemExit(f"Created {path.name}. Add product links, then run this script again.")


def read_urls(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit(f"{path.name} is empty")
    index = {header_key(col): i for i, col in enumerate(rows[0])}
    if "amazonurl" not in index:
        raise SystemExit("Urls.xlsx needs an amazon_url column")
    items: list[dict[str, str]] = []
    for values in rows[1:]:
        if not values:
            continue
        url = cell_text(values[index["amazonurl"]] if index["amazonurl"] < len(values) else "")
        if not url:
            continue
        def take(name: str) -> str:
            pos = index.get(name)
            if pos is None or pos >= len(values):
                return ""
            return cell_text(values[pos])

        status = take("status") or "Draft"
        shipping = take("freeshipping") or "False"
        items.append(
            {
                "amazon_url": url,
                "store_name": take("storename").upper(),
                "sku": take("sku"),
                "upc": take("upc"),
                "category": take("category"),
                "subcategory": take("subcategory"),
                "status": status,
                "free_shipping": shipping,
            }
        )
    if not items:
        raise SystemExit("Urls.xlsx has no product links")
    return items


def read_done_urls(path: Path, needs_again=None) -> set[str]:
    if not path.is_file():
        return set()
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return set()
    index = {header_key(col): i for i, col in enumerate(rows[0])}
    pos = index.get("amazonurl")
    if pos is None:
        return set()
    done: set[str] = set()
    for row in rows[1:]:
        if not row or pos >= len(row):
            continue
        url = cell_text(row[pos])
        if not url:
            continue
        if needs_again is not None and needs_again(row, index):
            continue
        done.add(url)
    return done


def append_create_row(path: Path, row: dict[str, str]) -> None:
    if path.is_file():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "create"
        ws.append(list(CREATE_COLUMNS))
    values = [row.get(col, "") for col in CREATE_COLUMNS]
    headers = [header_key(cell.value) for cell in ws[1]]
    url_col = next((i for i, key in enumerate(headers) if key == "amazonurl"), None)
    url = row.get("amazon_url", "")
    if url_col is not None and url:
        for excel_row in range(2, ws.max_row + 1):
            if cell_text(ws.cell(excel_row, url_col + 1).value) != url:
                continue
            for col, value in enumerate(values, start=1):
                ws.cell(excel_row, col).value = value
            wb.save(path)
            return
    ws.append(values)
    wb.save(path)


def start_browser(site: str = "Amazon", profile_dir: Optional[Path] = None) -> WebDriver:
    driver_path = local_chromedriver("")
    options = Options()
    options.add_argument("--disable-session-crashed-bubble")
    options.add_argument("--no-first-run")
    options.add_argument("--no-default-browser-check")
    if profile_dir is not None:
        profile_dir.mkdir(parents=True, exist_ok=True)
        options.add_argument(f"--user-data-dir={profile_dir}")
    log.info("Opening Chrome for %s with ChromeDriver %s", site, driver_path.parent.name)
    return webdriver.Chrome(service=Service(executable_path=str(driver_path)), options=options)


def page_blocked(driver: WebDriver) -> bool:
    body = ""
    try:
        body = driver.find_element(By.TAG_NAME, "body").text
    except Exception:
        return False
    return bool(re.search(r"enter the characters|not a robot|sorry, we just need to make sure", body, re.I))


def wait_if_blocked(driver: WebDriver) -> None:
    if not page_blocked(driver):
        return
    log.warning("Amazon asked for a check. Finish it in the Chrome window.")
    deadline = time.time() + 180
    while time.time() < deadline:
        if not page_blocked(driver):
            log.info("Amazon check cleared")
            return
        time.sleep(2)
    raise RuntimeError("Amazon check was not finished")


def set_zip(driver: WebDriver) -> None:
    try:
        current = driver.find_element(By.ID, "glow-ingress-line2").text
    except Exception:
        current = ""
    if ZIP_CODE in current:
        log.info("Amazon location already %s", ZIP_CODE)
        return
    log.info("Setting Amazon delivery ZIP to %s", ZIP_CODE)
    driver.find_element(By.ID, "nav-global-location-popover-link").click()
    box = WebDriverWait(driver, 10).until(EC.element_to_be_clickable((By.ID, "GLUXZipUpdateInput")))
    box.click()
    box.send_keys(Keys.CONTROL, "a")
    box.send_keys(ZIP_CODE)
    box.send_keys(Keys.ENTER)
    time.sleep(1.5)
    for selector in ("button[name='glowDoneButton']", "#GLUXConfirmClose", "input#GLUXConfirmClose"):
        buttons = driver.find_elements(By.CSS_SELECTOR, selector)
        if buttons:
            try:
                buttons[0].click()
                break
            except Exception:
                continue
    time.sleep(2)


def text_of(driver: WebDriver, selector: str) -> str:
    try:
        return re.sub(r"\s+", " ", driver.find_element(By.CSS_SELECTOR, selector).text).strip()
    except Exception:
        return ""


def clean_brand(raw: str) -> str:
    text = re.sub(r"\s+", " ", raw or "").strip()
    text = re.sub(r"(?i)^visit the\s+", "", text)
    text = re.sub(r"(?i)^brand:\s*", "", text)
    text = re.sub(r"(?i)\s+store$", "", text)
    return text.strip()


def money(raw: str) -> str:
    match = re.search(r"(\d[\d,]*(?:\.\d{2})?)", raw or "")
    if not match:
        return ""
    return match.group(1).replace(",", "")


def bullets_of(driver: WebDriver) -> list[str]:
    try:
        items = driver.execute_script(
            """
            const nodes = [...document.querySelectorAll('#feature-bullets li span.a-list-item')];
            return nodes.map(n => (n.innerText || '').trim()).filter(Boolean);
            """
        )
    except Exception:
        items = []
    return [re.sub(r"\s+", " ", str(item)).strip() for item in items or [] if str(item).strip()]


MIN_DESCRIPTION_POINTS = 4
_SKIP_FACT_NAMES = {
    "brand name",
    "choice",
    "origin",
    "place of origin",
    "ships from",
    "high-concerned chemical",
    "cn",
}


def _tidy_point(point: str) -> str:
    point = re.sub(r"\s+", " ", point or "").strip(" .-")
    if not point:
        return ""
    return point[0].upper() + point[1:]


def _point_words(point: str) -> set[str]:
    return set(re.findall(r"[a-z0-9]+", point.lower()))


def _point_from_fact(fact: str) -> str:
    name, sep, value = fact.partition(":")
    if not sep:
        return _tidy_point(fact)
    name = re.sub(r"\s+", " ", name).strip()
    value = re.sub(r"\s+", " ", value).strip()
    label = name.lower()
    if not name or label in _SKIP_FACT_NAMES:
        return ""
    if value.lower() in {"none", "no", "n/a", "false", "other"}:
        return ""
    if value.lower() in {"yes", "true"}:
        return _tidy_point(label)
    if "material" in label:
        kind = re.sub(r"\bmaterial\b", "", label).strip()
        return _tidy_point(f"{value} {kind}".strip())
    if " " in value or len(value) > 12:
        return _tidy_point(value)
    return _tidy_point(f"{value} {label}".strip())


def listing_description(title: str, facts: list[str]) -> str:
    points: list[str] = []
    seen: list[set[str]] = []

    def add(point: str) -> None:
        point = _tidy_point(point)
        words = _point_words(point)
        if len(point) < 2 or not words:
            return
        for existing in seen:
            if words <= existing or existing <= words:
                return
        seen.append(words)
        points.append(point)

    add(title)
    for fact in facts:
        add(_point_from_fact(fact))
    for extra in (
        "Brand new",
        "See the photos for the finish and what is included",
        "Packed and ready to ship",
    ):
        if len(points) >= MIN_DESCRIPTION_POINTS:
            break
        add(extra)
    return "\n".join(f"- {point}" for point in points)


def description_from_bullets(bullets: list[str]) -> str:
    return listing_description("", bullets)


def asin_of(url: str, page: str) -> str:
    match = ASIN_RE.search(url) or re.search(r"\bASIN\s*[:\s]+([A-Z0-9]{10})\b", page, re.I)
    return match.group(1).upper() if match else ""


def upc_of(page: str) -> str:
    match = re.search(r"\b(?:UPC|EAN|GTIN)\s*[:\s#]+(\d{8,14})\b", page, re.I)
    return match.group(1) if match else ""


def image_urls(page: str) -> list[str]:
    found = HIRES_RE.findall(page) or LARGE_RE.findall(page)
    urls: list[str] = []
    for url in found:
        if "sprite" in url or "transparent-pixel" in url or url in urls:
            continue
        urls.append(url.replace("\\u0026", "&"))
    return urls


_FACE_DETECTOR = None


def has_person(image: Image.Image) -> bool:
    model = ROOT / "models" / "face_detection_yunet_2023mar.onnx"
    if not model.is_file():
        return False
    rgb = image.convert("RGB")
    width, height = rgb.size
    scale = min(1.0, 640 / max(width, height))
    if scale < 1:
        rgb = rgb.resize((max(1, int(width * scale)), max(1, int(height * scale))), Image.Resampling.LANCZOS)
    frame = cv2.cvtColor(np.array(rgb), cv2.COLOR_RGB2BGR)
    global _FACE_DETECTOR
    if _FACE_DETECTOR is None:
        _FACE_DETECTOR = cv2.FaceDetectorYN.create(str(model), "", (frame.shape[1], frame.shape[0]), 0.8, 0.3, 5000)
    _FACE_DETECTOR.setInputSize((frame.shape[1], frame.shape[0]))
    try:
        _, faces = _FACE_DETECTOR.detect(frame)
    except cv2.error:
        return False
    if faces is None:
        return False
    return any(float(face[-1]) >= 0.7 for face in faces)


def _edge_has_logo(strip: np.ndarray) -> bool:
    if strip.size == 0:
        return False
    if float(strip.mean()) > 40:
        return False
    profile = strip.mean(axis=0)
    busy = profile > 30
    quiet = profile < 8
    return bool(busy.any() and busy.mean() <= 0.45 and quiet.mean() >= 0.4)


def crop_edge_branding(image: Image.Image) -> Image.Image:
    rgb = np.array(image.convert("RGB"))
    height, width = rgb.shape[:2]
    if width < 200 or height < 200:
        return image
    edges = cv2.Canny(cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY), 80, 160)
    top = max(8, int(height * 0.14))
    side = max(8, int(width * 0.14))
    y0, y1, x0, x1 = 0, height, 0, width
    if _edge_has_logo(edges[:top, :]):
        y0 = top
    if _edge_has_logo(edges[height - top :, :]):
        y1 = height - top
    if _edge_has_logo(edges[:, :side].T):
        x0 = side
    if _edge_has_logo(edges[:, width - side :].T):
        x1 = width - side
    if (x1 - x0) < width * 0.7 or (y1 - y0) < height * 0.7:
        return image
    if (x0, y0, x1, y1) == (0, 0, width, height):
        return image
    return image.crop((x0, y0, x1, y1))


def save_product_image(data: bytes, dest: Path) -> bool:
    try:
        image = Image.open(io.BytesIO(data))
    except Exception:
        return False
    image = image.convert("RGB")
    if image.width < 200 or image.height < 200:
        return False
    if has_person(image):
        log.info("Skipped a photo because a person is in the frame")
        return False
    image = crop_edge_branding(image)
    image.thumbnail((IMAGE_SIZE, IMAGE_SIZE), Image.Resampling.LANCZOS)
    dest.parent.mkdir(parents=True, exist_ok=True)
    image.save(dest, "JPEG", quality=90)
    return True


def image_url_fallbacks(url: str) -> list[str]:
    urls = [url]
    match = re.search(r"(/kf/[^?#]+)", url, re.I)
    if not match or not re.search(r"aliexpress-media\.com|alicdn\.com", url, re.I):
        return urls
    path = match.group(1)
    for host in ("https://ae01.alicdn.com", "https://img.alicdn.com"):
        alt = host + path
        if alt not in urls:
            urls.append(alt)
    return urls


def download_image(url: str) -> bytes:
    last_error: Optional[Exception] = None
    for candidate in image_url_fallbacks(url):
        for attempt in range(2):
            try:
                resp = requests.get(candidate, timeout=30, headers={"User-Agent": "Mozilla/5.0"})
                resp.raise_for_status()
                if resp.content:
                    return resp.content
            except Exception as exc:
                last_error = exc
                time.sleep(0.5)
    if last_error:
        raise last_error
    raise RuntimeError("Image download returned nothing")


def collect_photos(urls: list[str], sku: str) -> list[str]:
    saved: list[str] = []
    bunny = load_bunny()
    for url in urls:
        if len(saved) >= MAX_IMAGES:
            break
        try:
            content = download_image(url)
        except Exception as exc:
            log.warning("Image download failed: %s", exc)
            continue
        dest = PHOTO_DIR / f"{sku}_{len(saved) + 1}.jpg"
        if not save_product_image(content, dest):
            continue
        cdn = upload_photo(bunny, dest) if bunny else ""
        saved.append(cdn or str(dest))
    return saved


def load_bunny() -> dict[str, str]:
    candidates = [
        ROOT / ".env",
        Path(r"G:\All Tools Created\BunnyCDN IMAGES CONVERTER\.env"),
    ]
    values: dict[str, str] = {}
    for env_path in candidates:
        if not env_path.is_file():
            continue
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            values.setdefault(key.strip(), value.strip().strip('"').strip("'"))
        if values.get("BUNNY_STORAGE_ZONE") and values.get("BUNNY_STORAGE_API_KEY") and values.get("BUNNY_CDN_URL"):
            break
    zone = values.get("BUNNY_STORAGE_ZONE", "")
    key = values.get("BUNNY_STORAGE_API_KEY", "")
    pull = values.get("BUNNY_CDN_URL", "").rstrip("/")
    if not zone or not key or not pull:
        log.warning("Bunny settings were not found, photos will stay as local files")
        return {}
    region = values.get("BUNNY_STORAGE_REGION", "").strip().lower()
    endpoint = "https://storage.bunnycdn.com" if region in {"", "de", "falkenstein", "storage"} else f"https://{region}.storage.bunnycdn.com"
    prefix = values.get("BUNNY_FOLDER", "reverb").strip("/") or "reverb"
    if not (ROOT / ".env").is_file():
        prefix = "reverb"
    return {
        "zone": zone,
        "key": key,
        "pull": pull,
        "endpoint": endpoint,
        "prefix": prefix or "reverb",
    }


def upload_photo(cfg: dict[str, str], dest: Path) -> str:
    storage_path = "/".join(part for part in (cfg["prefix"], dest.name) if part)
    url = f"{cfg['endpoint']}/{cfg['zone']}/{storage_path}"
    try:
        response = requests.put(
            url,
            headers={"AccessKey": cfg["key"], "Content-Type": "image/jpeg"},
            data=dest.read_bytes(),
            timeout=90,
        )
    except requests.RequestException as exc:
        log.warning("Bunny upload failed for %s: %s", dest.name, exc)
        return ""
    if response.status_code not in {200, 201}:
        log.warning("Bunny upload failed for %s: HTTP %s", dest.name, response.status_code)
        return ""
    encoded = "/".join(quote(part, safe=".-_") for part in storage_path.split("/"))
    cdn = f"{cfg['pull']}/{encoded}"
    log.info("Uploaded %s", cdn)
    return cdn


def listing_quantity(stock: int) -> str:
    if stock <= 7:
        return "0"
    if stock <= 10:
        return "1"
    return "2"


def amazon_stock(driver: WebDriver) -> int:
    text = text_of(driver, "#availability") or text_of(driver, "#availability span")
    low = text.lower()
    if re.search(r"out of stock|currently unavailable|cannot be shipped", low):
        return 0
    left = re.search(r"only\s+(\d+)\s+left", low)
    if left:
        return int(left.group(1))
    try:
        options = driver.execute_script(
            """
            const sel = document.querySelector('#quantity, select[name="quantity"]');
            if (!sel) return [];
            return [...sel.options].map(o => parseInt(o.value, 10)).filter(n => !Number.isNaN(n));
            """
        )
    except Exception:
        options = []
    if options:
        return max(int(value) for value in options)
    if "in stock" in low:
        return 11
    return 0


_REVERB_TREE: Optional[tuple[dict[str, str], dict[str, list[tuple[str, str]]]]] = None


def load_reverb_categories() -> tuple[dict[str, str], dict[str, list[tuple[str, str]]]]:
    global _REVERB_TREE
    if _REVERB_TREE is not None:
        return _REVERB_TREE
    parents: dict[str, str] = {}
    children: dict[str, list[tuple[str, str]]] = {}
    candidates = [
        ROOT / "reverb_categories_cache.json",
        Path(r"G:\All Tools Created\Bulk Listing Formatter\output\reverb_categories_cache.json"),
        Path.home() / "Downloads" / "category uuid.json",
    ]
    for path in candidates:
        if not path.is_file():
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            log.warning("Could not read categories from %s: %s", path, exc)
            continue
        items = data.get("categories") if isinstance(data, dict) else data
        for item in items or []:
            if not isinstance(item, dict):
                continue
            full = str(item.get("full_name") or item.get("name") or "").strip()
            parts = [part.strip() for part in full.split("/") if part.strip()]
            if not parts:
                continue
            parents.setdefault(parts[0].lower(), parts[0])
            if len(parts) > 1:
                children.setdefault(parts[-1].lower(), []).append((parts[0], parts[-1]))
        if parents:
            log.info("Loaded %s Reverb categories from %s", len(parents), path.name)
            break
    _REVERB_TREE = (parents, children)
    return _REVERB_TREE


def parent_aliases(name: str) -> list[str]:
    names = [name]
    low = name.lower()
    if low.endswith(" kits"):
        stem = name[: -len(" kits")].strip()
        names.extend([f"{stem}s", stem])
    elif low.endswith(" kit"):
        stem = name[: -len(" kit")].strip()
        names.extend([f"{stem}s", stem])
    return names


def choose_child(hits: list[tuple[str, str]], earlier: list[str]) -> tuple[str, str]:
    unique = list(dict.fromkeys(hits))
    if len(unique) == 1:
        return unique[0]
    earlier_low = [name.lower() for name in earlier]
    exact = [hit for hit in unique if hit[0].lower() in earlier_low]
    if len(exact) == 1:
        return exact[0]
    trail = " ".join(earlier_low)
    if any(word in trail for word in ("musical instrument", "live sound", "stage", "studio", "guitar")):
        accessories = [hit for hit in unique if hit[0].lower() == "accessories"]
        if len(accessories) == 1:
            return accessories[0]
    if "home audio" in trail or "home theater" in trail:
        home = [hit for hit in unique if hit[0].lower() == "home audio"]
        if len(home) == 1:
            return home[0]
    return "", ""


def contained_children(crumb: str, children: dict[str, list[tuple[str, str]]]) -> list[tuple[str, str]]:
    words = re.findall(r"[a-z0-9]+", crumb.lower())
    last = words[-1] if words else ""
    found: list[tuple[str, str]] = []
    for label, hits in children.items():
        if len(label) < 5:
            continue
        label_words = re.findall(r"[a-z0-9]+", label)
        if not label_words:
            continue
        matched = False
        width = len(label_words)
        for start in range(0, len(words) - width + 1):
            window = words[start : start + width]
            if window == label_words or [part.rstrip("s") for part in window] == [part.rstrip("s") for part in label_words]:
                matched = True
                break
        if matched:
            found.extend(hits)
    ending = [hit for hit in found if re.findall(r"[a-z0-9]+", hit[1].lower())[-1].rstrip("s") == last.rstrip("s")]
    return ending or found


def map_reverb_category(crumbs: list[str]) -> tuple[str, str]:
    parents, children = load_reverb_categories()
    for index in range(len(crumbs) - 1, -1, -1):
        name = crumbs[index].strip()
        if not name:
            continue
        earlier = crumbs[:index]
        parent = parents.get(name.lower(), "")
        if parent:
            return parent, ""
        chosen = choose_child(children.get(name.lower(), []), earlier)
        if chosen[0]:
            return chosen
        for alias in parent_aliases(name):
            if alias.lower() == name.lower():
                continue
            parent = parents.get(alias.lower(), "")
            if parent:
                return parent, ""
        chosen = choose_child(contained_children(name, children), earlier)
        if chosen[0]:
            return chosen
        chosen = choose_child(last_word_children(name, children), earlier)
        if chosen[0]:
            return chosen
    return "", ""


def last_word_children(crumb: str, children: dict[str, list[tuple[str, str]]]) -> list[tuple[str, str]]:
    words = re.findall(r"[a-z0-9]+", crumb.lower())
    if not words:
        return []
    last = words[-1].rstrip("s")
    if len(last) < 5:
        return []
    found: list[tuple[str, str]] = []
    for label, hits in children.items():
        label_words = re.findall(r"[a-z0-9]+", label)
        if label_words and label_words[-1].rstrip("s") == last:
            found.extend(hits)
    return found


def amazon_breadcrumbs(driver: WebDriver) -> list[str]:
    try:
        crumbs = driver.execute_script(
            """
            const root = document.querySelector('#wayfinding-breadcrumbs_feature_div');
            if (!root) return [];
            return [...root.querySelectorAll('li')].map(li => (li.innerText || '').replace(/[›>]/g, ' ').replace(/\\s+/g, ' ').trim()).filter(Boolean);
            """
        )
    except Exception:
        crumbs = []
    return [str(crumb).strip() for crumb in crumbs or [] if str(crumb).strip()]


def scrape_one(driver: WebDriver, item: dict[str, str]) -> dict[str, str]:
    url = item["amazon_url"]
    if not re.search(r"https?://(?:www\.)?amazon\.com/", url, re.I):
        raise RuntimeError("Only amazon.com links are scraped. Use scrape_AliexpressAU.py for AliExpress.")
    log.info("Scraping %s", url)
    driver.get(url)
    time.sleep(2)
    wait_if_blocked(driver)
    set_zip(driver)
    driver.get(url)
    time.sleep(2)
    wait_if_blocked(driver)
    WebDriverWait(driver, 20).until(EC.presence_of_element_located((By.ID, "productTitle")))
    page = driver.page_source
    title = text_of(driver, "#productTitle") or text_of(driver, "span.product-title-word-break")
    if not title:
        raise RuntimeError("Product title was not found")
    price = money(text_of(driver, ".a-price .a-offscreen") or text_of(driver, ".a-price"))
    brand = clean_brand(text_of(driver, "#bylineInfo"))
    sku = item.get("sku") or asin_of(url, driver.find_element(By.TAG_NAME, "body").text)
    upc = item.get("upc") or upc_of(driver.find_element(By.TAG_NAME, "body").text)
    stock = amazon_stock(driver)
    quantity = listing_quantity(stock)
    log.info("Amazon stock %s mapped to quantity %s", stock, quantity)
    crumbs = amazon_breadcrumbs(driver)
    category, subcategory = map_reverb_category(crumbs)
    if category:
        log.info("Amazon categories %s -> %s / %s", " > ".join(crumbs), category, subcategory or "(none)")
    else:
        category = item.get("category", "")
        subcategory = item.get("subcategory", "")
        log.warning("No Reverb category for %s", " > ".join(crumbs) or "(no breadcrumb)")
    photos = collect_photos(image_urls(page), sku or "photo")
    log.info("Kept %s photos", len(photos))
    return {
        "store_name": item["store_name"],
        "title": title,
        "description": listing_description(title, bullets_of(driver)) or title,
        "remove_branding": brand,
        "price": price,
        "quantity": quantity,
        "Category": category,
        "Subcategory": subcategory,
        "sku": sku,
        "photos": "\n".join(photos),
        "status": item["status"] or "Draft",
        "upc": upc,
        "free_shipping": item["free_shipping"] or "False",
        "amazon_url": url,
    }


def main() -> None:
    ensure_urls_file()
    items = read_urls(URLS_PATH)
    amazon = [item for item in items if re.search(r"https?://(?:www\.)?amazon\.com/", item["amazon_url"], re.I)]
    other = len(items) - len(amazon)
    if other:
        log.info("Skipping %s link(s) that are not amazon.com. Those belong in scrape_aliexpress/Urls.xlsx", other)
    if not amazon:
        raise SystemExit("Urls.xlsx has no amazon.com links")
    done = read_done_urls(CREATE_PATH)
    pending = [item for item in amazon if item["amazon_url"] not in done]
    if not pending:
        log.info("%s already has every Amazon link. Check it, then run create_listing/create_listing.py", CREATE_PATH.name)
        return
    driver: Optional[WebDriver] = None
    failed: list[str] = []
    try:
        driver = start_browser()
        for item in pending:
            try:
                append_create_row(CREATE_PATH, scrape_one(driver, item))
                log.info("Wrote %s", item["amazon_url"])
            except Exception as exc:
                failed.append(item["amazon_url"])
                log.error("Failed %s: %s", item["amazon_url"], exc)
    finally:
        if driver is not None:
            driver.quit()
    if failed:
        raise SystemExit(f"{len(failed)} link(s) failed. Fix them and run the scraper again.")
    log.info("Create.xlsx is complete. Check the file, then run: python create_listing/create_listing.py create_listing/Create.xlsx")


if __name__ == "__main__":
    main()
