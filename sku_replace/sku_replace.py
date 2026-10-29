#!/usr/bin/env python3
"""
Replace the SKU (and optional UPC) on existing Reverb listings.

Excel columns:
  store_name, sku, new_sku, upc

sku is the SKU already on the listing. A blank upc cell is left unchanged.

Usage:
  python sku_replace/sku_replace.py sku_replace/SkuReplace.xlsx
  python sku_replace/sku_replace.py --create-sample
  python sku_replace/sku_replace.py sku_replace/SkuReplace.xlsx --store GG --no-vpn --skip-ip-check
"""
from __future__ import annotations

import argparse
import logging
import sys
from dataclasses import dataclass
from pathlib import Path

BOT_DIR = Path(__file__).resolve().parent
PROJECT = BOT_DIR.parent
LOG_PATH = BOT_DIR / "sku_replace.log"
DEFAULT_EXCEL = BOT_DIR / "SkuReplace.xlsx"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(LOG_PATH, encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger("sku_replace")

sys.path.insert(0, str(PROJECT / "shared"))
sys.path.insert(0, str(PROJECT / "create_listing"))

from openpyxl import Workbook, load_workbook

import create_listing as cl
import reverb_edit


HEADER_MAP = {
    "store name": "store",
    "store_name": "store",
    "storename": "store",
    "store": "store",
    "sku": "sku",
    "old sku": "sku",
    "old_sku": "sku",
    "new sku": "new_sku",
    "new_sku": "new_sku",
    "newsku": "new_sku",
    "upc": "upc",
    "barcode": "upc",
    "ean": "upc",
}


@dataclass
class SkuRow:
    row_number: int
    store: str
    sku: str
    new_sku: str
    upc: str
    status: str = ""
    detail: str = ""


def create_sample(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "input"
    ws.append(["store_name", "sku", "new_sku", "upc"])
    ws.append(["GG", "OLD-SKU-001", "NEW-SKU-001", ""])
    for cell in ws[2]:
        cell.number_format = "@"
    wb.save(path)
    log.info("Wrote sample Excel: %s", path)


def read_excel(path: Path) -> list[SkuRow]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    header_cells = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_cells:
        raise SystemExit("Excel has no header row")
    index: dict[str, int] = {}
    for i, col in enumerate(header_cells):
        mapped = HEADER_MAP.get(str(col or "").strip().lower())
        if mapped and mapped not in index:
            index[mapped] = i
    missing = [name for name in ("store", "sku", "new_sku") if name not in index]
    if missing:
        raise SystemExit(f"Excel is missing columns: {', '.join(missing)}")

    def cell(values: tuple, key: str):
        pos = index.get(key)
        if pos is None or pos >= len(values):
            return None
        return values[pos]

    rows: list[SkuRow] = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not values:
            continue
        store = cl.coerce_code(cell(values, "store")).upper()
        sku = cl.coerce_code(cell(values, "sku"))
        new_sku = cl.coerce_code(cell(values, "new_sku"))
        upc = cl.coerce_code(cell(values, "upc")) if "upc" in index else ""
        if not store and not sku and not new_sku:
            continue
        row = SkuRow(row_number=row_number, store=store, sku=sku, new_sku=new_sku, upc=upc)
        if not sku:
            row.status = "FAILED"
            row.detail = "SKU is blank"
        elif not new_sku:
            row.status = "FAILED"
            row.detail = "NEW SKU is blank"
        rows.append(row)
    return rows


def apply_row(driver, row: SkuRow) -> None:
    reverb_edit.update_sku(driver, row.sku, row.new_sku, row.upc)


def result_line(row: SkuRow) -> list:
    return [row.store, row.sku, row.new_sku, row.upc, row.status, row.detail]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Replace Reverb listing SKU and optional UPC from Excel")
    parser.add_argument("excel", nargs="?", default=str(DEFAULT_EXCEL), help="Excel file path")
    parser.add_argument("--store", help="Process only this store name")
    parser.add_argument("--limit", type=int, default=0, help="Max rows per store (0 = all)")
    parser.add_argument("--create-sample", action="store_true", help="Write SkuReplace.xlsx and exit")
    parser.add_argument("--no-vpn", action="store_true")
    parser.add_argument("--skip-ip-check", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.create_sample:
        create_sample(DEFAULT_EXCEL)
        return 0
    excel_path = Path(args.excel)
    if not excel_path.is_file():
        log.error("Excel not found: %s", excel_path)
        log.error("Create one with: python sku_replace/sku_replace.py --create-sample")
        return 1

    cfg = cl._load_config()
    chrome, user_data, vpncmd = cl.resolve_paths(cfg)
    rows = read_excel(excel_path)
    if args.store:
        only = args.store.strip().upper()
        rows = [row for row in rows if row.store == only]
    if not rows:
        log.error("No rows to process")
        return 1

    grouped: dict[str, list[SkuRow]] = {}
    for row in rows:
        grouped.setdefault(row.store, []).append(row)
    if args.limit and args.limit > 0:
        grouped = {key: value[: args.limit] for key, value in grouped.items()}
        rows = [row for group in grouped.values() for row in group]

    log.info("Loaded %s SKU rows across %s store(s)", len(rows), len(grouped))
    results = reverb_edit.result_path("sku_replace_results")
    headers = ["STORE NAME", "SKU", "NEW SKU", "UPC", "STATUS", "DETAIL"]

    for store_name, store_rows in grouped.items():
        info = cl.store_lookup(cfg, store_name)
        pending = [row for row in store_rows if not row.status]
        if not info:
            for row in pending:
                row.status = "FAILED"
                row.detail = f"Unknown store '{store_name}'"
        elif pending:
            reverb_edit.run_store_rows(
                cfg,
                chrome,
                user_data,
                vpncmd,
                info,
                store_rows,
                apply_row,
                use_vpn=not args.no_vpn,
                check_ip=not args.skip_ip_check,
            )
        reverb_edit.write_sheet(results, headers, [result_line(row) for row in rows])

    reverb_edit.write_sheet(results, headers, [result_line(row) for row in rows])
    ok = sum(1 for row in rows if row.status == "OK")
    failed = sum(1 for row in rows if row.status == "FAILED")
    log.info("Done. OK=%s  FAILED=%s", ok, failed)
    log.info("Results: %s", results)
    return 0 if failed == 0 else 2


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        log.error("Stopped by user")
        raise SystemExit(130)
